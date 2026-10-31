import os
import asyncio
import json
import base64
import uuid
import time
import re
import html
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from flow import (
    create_requirement_analysis_flow,
    create_exec_case_flow,
    create_system_test_flow,
    create_test_case_flow,
    create_test_design_flow,
)
from schemas import DEFAULT_ACTION_VOCABULARY, DEFAULT_ASSERTION_VOCABULARY, DEFAULT_CAPABILITIES
from utils.bugzilla import BugzillaSettings, bugs_to_requirement_text, fetch_bugs
from utils.llm_client import LLMClient, LLMSettings
from utils.paths import CONFIG_DIR, HISTORY_DIR, PROMPT_OVERRIDE_DIR, RAG_DIR, STATIC_DIR, prompt_path
from utils.rag_store import SimpleRAGStore


class PromptSettingsRequest(BaseModel):
    prompts: Dict[str, str]


class ExportTestCasesRequest(BaseModel):
    requirement_input: Optional[List[str]] = None
    requirement_spec: Optional[Dict[str, Any]] = None
    persona_reviews: Optional[Dict[str, Any]] = None
    requirement_review_history: Optional[Dict[str, Any]] = None
    test_design_spec: Optional[Dict[str, Any]] = None
    test_case_spec: Dict[str, Any]
    executable_cases: Optional[Dict[str, Any]] = None


class RAGUploadBase64Request(BaseModel):
    filename: str
    content_base64: str


class ChatTurn(BaseModel):
    role: str
    content: str


class ChatRequest(BaseModel):
    message: str
    history: List[ChatTurn] = Field(default_factory=list)


class RequirementHITLStartRequest(BaseModel):
    requirements: List[str] = Field(..., min_length=1)
    rag_context: str = ""
    framework_capability_catalog: str = ""
    product_profile: Optional[Dict[str, Any]] = None
    test_environments: Optional[Dict[str, Any]] = None


class RequirementHITLNextRequest(BaseModel):
    open_question_answers: Optional[Dict[str, str]] = None


class StageGenerateRequest(BaseModel):
    requirement_input: Optional[List[str]] = None
    requirement_spec: Dict[str, Any]
    persona_reviews: Optional[Dict[str, Any]] = None
    requirement_review_history: Optional[Dict[str, Any]] = None
    test_design_spec: Optional[Dict[str, Any]] = None
    test_case_spec: Optional[Dict[str, Any]] = None
    product_profile: Optional[Dict[str, Any]] = None
    test_environments: Optional[Dict[str, Any]] = None
    capabilities: Optional[Dict[str, Any]] = None
    action_vocabulary: Optional[List[str]] = None
    assertion_vocabulary: Optional[List[str]] = None
    history_record_id: str = ""
    bug_ids: Optional[List[str]] = None
    route_overrides: Optional[Dict[str, str]] = None


class BugzillaSettingsRequest(BaseModel):
    url: str = Field(..., min_length=1)
    api_key: str = ""
    clear_key: bool = False


class LLMSettingsRequest(BaseModel):
    api_base: str = ""
    model: str = ""
    api_key: str = ""
    clear_key: bool = False
    timeout: Optional[int] = None


class BugzillaFetchRequest(BaseModel):
    ids: str = Field(..., min_length=1)
    excluded_authors: str = ""


app = FastAPI(title="TestFlow")
static_dir = str(STATIC_DIR)
design_history_dir = HISTORY_DIR
rag_store = SimpleRAGStore(RAG_DIR)
bugzilla_settings = BugzillaSettings(CONFIG_DIR / "bugzilla-settings.json")
requirement_hitl_jobs: Dict[str, Dict[str, Any]] = {}
stage_jobs: Dict[str, Dict[str, Any]] = {}
ALLOWED_PROMPT_FILES = {
    "req_parse.txt",
    "review_spec_lawyer.txt",
    "review_carrier.txt",
    "review_ux_advocate.txt",
    "synthesis_arbiter.txt",
    "test_designer.txt",
    "integrated_matrix_designer.txt",
    "test_case_generator.txt",
    "exec_case_artemis.txt",
    "exec_case_at_agent.txt",
    "exec_case_amarisoft.txt",
    "exec_case_protocol.txt",
}
app.mount("/static", StaticFiles(directory=static_dir), name="static")


design_history_dir.mkdir(parents=True, exist_ok=True)


def _valid_history_record_id(record_id: str) -> bool:
    return bool(re.fullmatch(r"[A-Za-z0-9_-]{6,64}", record_id or ""))


def _history_record_dir(record_id: str) -> Path:
    if not _valid_history_record_id(record_id):
        raise HTTPException(status_code=404, detail="history记录不存在")
    return design_history_dir / record_id


def _history_record_json_path(record_id: str) -> Path:
    return _history_record_dir(record_id) / "record.json"


def _history_document_path(record_id: str) -> Path:
    return _history_record_dir(record_id) / "design_document.html"


def _history_document_url(record_id: str) -> str:
    return f"/api/design-history/{record_id}/document"


def _safe_json_text(value: Any) -> str:
    try:
        return json.dumps(value, ensure_ascii=False, indent=2)
    except Exception:
        return str(value)


def _history_title(requirement_input: List[str], requirement_spec: Dict[str, Any]) -> str:
    reqs = requirement_spec.get("final_requirements", []) if isinstance(requirement_spec, dict) else []
    if isinstance(reqs, list):
        for item in reqs:
            if isinstance(item, dict) and str(item.get("title", "")).strip():
                return str(item.get("title")).strip()
    for raw in requirement_input or []:
        text = str(raw or "").strip()
        if text:
            return text[:80]
    return "未命名测试设计"


def _history_meta_counts(record: Dict[str, Any]) -> Dict[str, int]:
    requirement_spec = record.get("requirement_spec", {}) if isinstance(record, dict) else {}
    test_design_spec = record.get("test_design_spec", {}) if isinstance(record, dict) else {}
    test_case_spec = record.get("test_case_spec", {}) if isinstance(record, dict) else {}
    return {
        "requirements": len(requirement_spec.get("final_requirements", []) if isinstance(requirement_spec, dict) else []),
        "objectives": len(test_design_spec.get("objectives", []) if isinstance(test_design_spec, dict) else []),
        "matrix_rows": len(test_design_spec.get("integrated_matrix", []) if isinstance(test_design_spec, dict) else []),
        "testcases": len(test_case_spec.get("testcases", []) if isinstance(test_case_spec, dict) else []),
    }


def _history_summary(record: Dict[str, Any]) -> Dict[str, Any]:
    counts = _history_meta_counts(record)
    record_id = str(record.get("record_id", ""))
    return {
        "record_id": record_id,
        "title": record.get("title", "未命名测试设计"),
        "status": record.get("status", "design_completed"),
        "created_at": record.get("created_at", ""),
        "updated_at": record.get("updated_at", ""),
        "requirement_count": counts["requirements"],
        "objective_count": counts["objectives"],
        "matrix_row_count": counts["matrix_rows"],
        "testcase_count": counts["testcases"],
        "has_testcases": counts["testcases"] > 0,
        "document_url": _history_document_url(record_id),
    }


def _render_html_list(items: Any) -> str:
    arr = [str(x).strip() for x in (items or []) if str(x).strip()]
    if not arr:
        return "<p>无</p>"
    return "<ul>" + "".join(f"<li>{html.escape(x)}</li>" for x in arr) + "</ul>"


def _render_html_kv_table(rows: List[List[str]]) -> str:
    if not rows:
        return "<p>无</p>"
    body = "".join(
        f"<tr><th>{html.escape(str(k))}</th><td>{html.escape(str(v))}</td></tr>"
        for k, v in rows
    )
    return f"<table><tbody>{body}</tbody></table>"


def _render_test_environment_lines(test_environments: Dict[str, Any]) -> List[str]:
    if not isinstance(test_environments, dict):
        return []
    display = test_environments.get("display")
    if isinstance(display, list):
        lines = [str(x).strip() for x in display if str(x).strip()]
        if lines:
            return lines
    details = test_environments.get("selected_details")
    if isinstance(details, list):
        lines = []
        for item in details:
            if not isinstance(item, dict):
                continue
            label = str(item.get("label", "")).strip()
            desc = str(item.get("description", "")).strip()
            priority = str(item.get("priority", "")).strip()
            prefix = f"{priority}. " if priority else ""
            if label and desc:
                lines.append(f"{prefix}{label}: {desc}")
            elif label:
                lines.append(f"{prefix}{label}")
        if lines:
            return lines
    return []


def _format_test_environment_cell(env: Any) -> tuple[str, str]:
    if not isinstance(env, dict):
        return ("未指定", "")
    primary = str(env.get("primary_label") or env.get("primary") or "未指定").strip() or "未指定"
    alternatives = env.get("alternative_labels")
    if not isinstance(alternatives, list) or not alternatives:
        alternatives = env.get("alternatives") if isinstance(env.get("alternatives"), list) else []
    alt_text = "、".join(str(x).strip() for x in alternatives if str(x).strip())
    rationale = str(env.get("rationale", "")).strip()
    risk = str(env.get("availability_risk", "")).strip()

    parts = [primary]
    if alt_text:
        parts.append(f"备选: {alt_text}")
    if rationale:
        parts.append(f"依据: {rationale}")
    return ("；".join(parts), risk)


def _parse_answer_map(answer_map: Dict[str, Any], source_round: int) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    if not isinstance(answer_map, dict):
        return out
    for raw_key, raw_answer in answer_map.items():
        answer = str(raw_answer or "").strip()
        if not answer:
            continue
        key = str(raw_key or "")
        req_id, question = "", key
        if "::" in key:
            req_id, question = key.split("::", 1)
        out.append(
            {
                "source_round": source_round,
                "req_id": req_id.strip(),
                "question": question.strip(),
                "answer": answer,
            }
        )
    return out


def _merge_answered_questions(existing: Any, additions: Any) -> List[Dict[str, Any]]:
    merged: Dict[tuple[str, str], Dict[str, Any]] = {}
    for item in list(existing or []) + list(additions or []):
        if not isinstance(item, dict):
            continue
        req_id = str(item.get("req_id", "")).strip()
        question = str(item.get("question", "")).strip()
        answer = str(item.get("answer", "")).strip()
        if not question or not answer:
            continue
        merged[(req_id, question)] = {
            "source_round": int(item.get("source_round", 0) or 0),
            "req_id": req_id,
            "question": question,
            "answer": answer,
            "source": str(item.get("source") or "用户"),
        }
    return list(merged.values())


def _build_requirement_review_history(rounds: Any, answered_questions: Any) -> Dict[str, Any]:
    round_list = [item for item in (rounds or []) if isinstance(item, dict)]
    answered_list = [item for item in (answered_questions or []) if isinstance(item, dict)]
    return {
        "rounds": round_list,
        "answered_questions": answered_list,
        "answered_count": len(answered_list),
    }


def _render_history_review_table(items: List[Dict[str, Any]]) -> str:
    if not items:
        return "<p>无</p>"
    rows = []
    for it in items:
        rows.append(
            "<tr>"
            f"<td>{html.escape(str(it.get('req_id', '')))}</td>"
            f"<td>{html.escape('；'.join(str(x) for x in it.get('issues', []) or []))}</td>"
            f"<td>{html.escape('；'.join(str(x) for x in it.get('open_questions', []) or []))}</td>"
            f"<td>{html.escape(str(it.get('rewrite_suggestion', '')))}</td>"
            "</tr>"
        )
    return (
        "<table><thead><tr><th>需求ID</th><th>问题项</th><th>待确认问题</th><th>改写建议</th></tr></thead>"
        f"<tbody>{''.join(rows)}</tbody></table>"
    )


def _render_answered_question_table(items: Any) -> str:
    rows = [it for it in (items or []) if isinstance(it, dict)]
    if not rows:
        return "<p>无</p>"
    body = []
    for item in rows:
        body.append(
            "<tr>"
            f"<td>{html.escape(str(item.get('source_round', '')))}</td>"
            f"<td>{html.escape(str(item.get('req_id', '')))}</td>"
            f"<td>{html.escape(str(item.get('question', '')))}</td>"
            f"<td>{html.escape(str(item.get('answer', '')))}</td>"
            f"<td>{html.escape(str(item.get('source', '用户')))}</td>"
            "</tr>"
        )
    return (
        "<table><thead><tr><th>来源轮次</th><th>需求ID</th><th>已确认问题</th><th>确认结果</th><th>确认方</th></tr></thead>"
        f"<tbody>{''.join(body)}</tbody></table>"
    )


def _render_design_document_html(record: Dict[str, Any]) -> str:
    requirement_input = record.get("requirement_input", []) or []
    product_profile = record.get("product_profile", {}) if isinstance(record.get("product_profile"), dict) else {}
    test_environments = record.get("test_environments", {}) if isinstance(record.get("test_environments"), dict) else {}
    requirement_spec = record.get("requirement_spec", {}) if isinstance(record.get("requirement_spec"), dict) else {}
    persona_reviews = record.get("persona_reviews", {}) if isinstance(record.get("persona_reviews"), dict) else {}
    requirement_review_history = record.get("requirement_review_history", {}) if isinstance(record.get("requirement_review_history"), dict) else {}
    test_design_spec = record.get("test_design_spec", {}) if isinstance(record.get("test_design_spec"), dict) else {}
    reqs = requirement_spec.get("final_requirements", []) if isinstance(requirement_spec.get("final_requirements"), list) else []
    objectives = test_design_spec.get("objectives", []) if isinstance(test_design_spec.get("objectives"), list) else []
    matrices = test_design_spec.get("coverage_matrices", []) if isinstance(test_design_spec.get("coverage_matrices"), list) else []
    integrated = test_design_spec.get("integrated_matrix", []) if isinstance(test_design_spec.get("integrated_matrix"), list) else []
    title = record.get("title") or _history_title(requirement_input, requirement_spec)
    updated_at = record.get("updated_at", "")

    req_rows = []
    for item in reqs:
        req_rows.append(
            "<tr>"
            f"<td>{html.escape(str(item.get('req_id', '')))}</td>"
            f"<td>{html.escape(str(item.get('title', '')))}</td>"
            f"<td>{html.escape(str(item.get('priority', '')))}</td>"
            f"<td>{html.escape('、'.join(str(x) for x in item.get('rat_scope', []) or []))}</td>"
            f"<td>{html.escape(' / '.join(str(x) for x in ((item.get('acceptance') or {}).get('pass_fail', []) or [])))}</td>"
            "</tr>"
        )
    req_table = (
        "<table><thead><tr><th>需求ID</th><th>标题</th><th>优先级</th><th>RAT范围</th><th>通过标准</th></tr></thead>"
        f"<tbody>{''.join(req_rows)}</tbody></table>"
        if req_rows
        else "<p>无</p>"
    )

    review_sections = []
    for label, key in [("规范评审", "spec"), ("运营商评审", "carrier"), ("现网体验评审", "ux")]:
        reviews = (persona_reviews.get(key) or {}).get("reviews", []) if isinstance(persona_reviews.get(key), dict) else []
        review_sections.append(f"<section><h3>{html.escape(label)}</h3>{_render_history_review_table(reviews)}</section>")

    obj_rows = []
    for item in objectives:
        obj_rows.append(
            "<tr>"
            f"<td>{html.escape(str(item.get('objective_id', '')))}</td>"
            f"<td>{html.escape('、'.join(str(x) for x in item.get('linked_reqs', []) or []))}</td>"
            f"<td>{html.escape(str(item.get('goal', '')))}</td>"
            f"<td>{html.escape(' / '.join(str(x) for x in item.get('success_criteria', []) or []))}</td>"
            f"<td>{html.escape(' / '.join(str(x) for x in item.get('evidence', []) or []))}</td>"
            f"<td>{html.escape(str(item.get('priority', '')))}</td>"
            "</tr>"
        )
    objective_table = (
        "<table><thead><tr><th>目标ID</th><th>关联需求</th><th>目标</th><th>成功标准</th><th>证据</th><th>优先级</th></tr></thead>"
        f"<tbody>{''.join(obj_rows)}</tbody></table>"
        if obj_rows
        else "<p>无</p>"
    )

    matrix_cards = []
    for matrix in matrices:
        dims = matrix.get("dimensions", []) if isinstance(matrix, dict) else []
        dim_html = "<br/>".join(
            f"{html.escape(str(d.get('name', '')))}: {html.escape('、'.join(str(x) for x in (d.get('values', []) or [])))}"
            for d in dims if isinstance(d, dict)
        )
        sampling = matrix.get("sampling_strategy", {}) if isinstance(matrix, dict) else {}
        matrix_cards.append(
            "<article class='matrix-card'>"
            f"<h4>{html.escape(str(matrix.get('matrix_id', '覆盖矩阵')))}</h4>"
            f"<div>{dim_html or '无'}</div>"
            f"<p><strong>采样策略：</strong>{html.escape(str(sampling.get('type', '')))}</p>"
            f"<p><strong>必选组合：</strong>{html.escape(_safe_json_text(sampling.get('must_include', [])))}</p>"
            f"<p><strong>排除组合：</strong>{html.escape(_safe_json_text(sampling.get('exclude', [])))}</p>"
            "</article>"
        )
    matrix_html = "".join(matrix_cards) if matrix_cards else "<p>无</p>"

    integrated_rows = []
    for idx, row in enumerate(integrated, start=1):
        cfg = row.get("key_configuration", {}) if isinstance(row, dict) and isinstance(row.get("key_configuration"), dict) else {}
        cfg_text = " / ".join(f"{k}: {v}" for k, v in cfg.items())
        env_text, risk_text = _format_test_environment_cell(row.get("test_environment"))
        integrated_rows.append(
            "<tr>"
            f"<td>{html.escape(str(row.get('row_id', f'ROW-{idx:03d}')))}</td>"
            f"<td>{html.escape(str(row.get('req_id', '')))}</td>"
            f"<td>{html.escape(str(row.get('objective_id', '')))}</td>"
            f"<td>{html.escape(str(row.get('scenario', '')))}</td>"
            f"<td>{html.escape(cfg_text)}</td>"
            f"<td>{html.escape(' / '.join(str(x) for x in row.get('pass_criteria', []) or []))}</td>"
            f"<td>{html.escape(env_text)}</td>"
            f"<td>{html.escape(risk_text)}</td>"
            "</tr>"
        )
    integrated_table = (
        "<table><thead><tr><th>行ID</th><th>需求ID</th><th>关联目标</th><th>测试场景</th><th>配置选取</th><th>通过标准</th><th>推荐测试环境</th><th>资源风险</th></tr></thead>"
        f"<tbody>{''.join(integrated_rows)}</tbody></table>"
        if integrated_rows
        else "<p>无</p>"
    )

    profile_lines = []
    for item in product_profile.get("profile_display", []) if isinstance(product_profile.get("profile_display"), list) else []:
        if isinstance(item, dict):
            group = str(item.get("group", "")).strip()
            values = item.get("values", []) or []
            if group and values:
                profile_lines.append(f"{group}: {'、'.join(str(x) for x in values)}")
    test_environment_lines = _render_test_environment_lines(test_environments)
    answered_question_table = _render_answered_question_table(requirement_review_history.get("answered_questions", []))

    return f"""<!DOCTYPE html>
<html lang="zh-CN">
  <head>
    <meta charset="UTF-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1.0" />
    <title>{html.escape(str(title))} - 测试设计文档</title>
    <style>
      body {{ margin: 0; padding: 32px; font-family: "Segoe UI", "PingFang SC", "Hiragino Sans GB", "Microsoft YaHei", sans-serif; background: #f4efe5; color: #1f2937; }}
      .page {{ max-width: 1180px; margin: 0 auto; background: #fffdf8; border: 1px solid #e5dccd; border-radius: 18px; box-shadow: 0 18px 48px rgba(15, 23, 42, 0.08); overflow: hidden; }}
      .hero {{ padding: 28px 32px 24px; background: linear-gradient(135deg, #f6efe3 0%, #fffaf1 58%, #eef6f7 100%); border-bottom: 1px solid #e9ddcb; }}
      .hero h1 {{ margin: 0; font-size: 30px; }}
      .hero p {{ margin: 8px 0 0; color: #5b6472; }}
      .section {{ padding: 24px 32px; border-top: 1px solid #efe6d8; }}
      .section:first-of-type {{ border-top: none; }}
      h2 {{ margin: 0 0 14px; font-size: 20px; }}
      h3 {{ margin: 0 0 10px; font-size: 16px; }}
      h4 {{ margin: 0 0 8px; font-size: 14px; }}
      table {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
      th, td {{ border: 1px solid #e5dccd; padding: 9px 10px; text-align: left; vertical-align: top; }}
      th {{ background: #f8f1e7; }}
      ul {{ margin: 8px 0 0 18px; }}
      p {{ line-height: 1.7; margin: 8px 0; }}
      .grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 16px; }}
      .matrix-card {{ border: 1px solid #e7ddcf; border-radius: 12px; padding: 14px; background: #fff; margin-bottom: 12px; }}
      @media (max-width: 900px) {{ body {{ padding: 12px; }} .section, .hero {{ padding: 18px; }} .grid {{ grid-template-columns: 1fr; }} }}
    </style>
  </head>
  <body>
    <div class="page">
      <section class="hero">
        <h1>标准测试设计文档</h1>
        <p>{html.escape(str(title))}</p>
      </section>
      <section class="section">
        <h2>一、文档信息</h2>
        {_render_html_kv_table([["记录ID", record.get("record_id", "")], ["最近更新时间", updated_at], ["当前状态", record.get("status", "")]])}
      </section>
      <section class="section">
        <h2>二、需求输入与产品背景</h2>
        <div class="grid">
          <div><h3>原始需求输入</h3>{_render_html_list(requirement_input)}</div>
          <div><h3>产品通信背景与支持特性</h3>{_render_html_list(profile_lines)}</div>
          <div><h3>测试团队可用测试环境</h3>{_render_html_list(test_environment_lines)}</div>
        </div>
      </section>
      <section class="section"><h2>三、测试需求解读稿</h2>{req_table}</section>
      <section class="section"><h2>四、三人评审记录</h2>{''.join(review_sections)}<section><h3>用户已确认问题</h3>{answered_question_table}</section></section>
      <section class="section"><h2>五、测试目标</h2>{objective_table}</section>
      <section class="section"><h2>六、覆盖矩阵</h2>{matrix_html}</section>
      <section class="section"><h2>七、整合覆盖矩阵</h2>{integrated_table}</section>
      <section class="section"><h2>八、设计说明</h2>{_render_html_list(test_design_spec.get("design_notes", []))}</section>
      <section class="section">
        <h2>九、范围裁剪与待跟进项</h2>
        <div class="grid">
          <div><h3>去范围项</h3>{_render_html_list(test_design_spec.get("de_scoped", []))}</div>
          <div><h3>待确认问题</h3>{_render_html_list(requirement_spec.get("questions_to_ask", []))}</div>
        </div>
      </section>
    </div>
  </body>
</html>"""


def _write_design_history_record(record: Dict[str, Any]) -> Dict[str, Any]:
    record_id = str(record.get("record_id") or uuid.uuid4().hex[:12])
    now = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime())
    record["record_id"] = record_id
    record["title"] = _history_title(record.get("requirement_input", []) or [], record.get("requirement_spec", {}) or {})
    record["created_at"] = record.get("created_at") or now
    record["updated_at"] = now
    record.setdefault("status", "design_completed")
    record_dir = _history_record_dir(record_id)
    record_dir.mkdir(parents=True, exist_ok=True)
    _history_record_json_path(record_id).write_text(_safe_json_text(record), encoding="utf-8")
    _history_document_path(record_id).write_text(_render_design_document_html(record), encoding="utf-8")
    return record


def _load_design_history_record(record_id: str) -> Dict[str, Any]:
    path = _history_record_json_path(record_id)
    if not path.exists():
        raise HTTPException(status_code=404, detail="history记录不存在")
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"history记录损坏: {exc}")


def _upsert_design_history_record(
    *,
    record_id: str = "",
    requirement_input: Optional[List[str]] = None,
    requirement_spec: Optional[Dict[str, Any]] = None,
    persona_reviews: Optional[Dict[str, Any]] = None,
    requirement_review_history: Optional[Dict[str, Any]] = None,
    test_design_spec: Optional[Dict[str, Any]] = None,
    product_profile: Optional[Dict[str, Any]] = None,
    test_environments: Optional[Dict[str, Any]] = None,
    test_case_spec: Optional[Dict[str, Any]] = None,
    warnings: Optional[Dict[str, Any]] = None,
    trace: Optional[Dict[str, Any]] = None,
    status: Optional[str] = None,
) -> Dict[str, Any]:
    record: Dict[str, Any] = {}
    if record_id:
        try:
            record = _load_design_history_record(record_id)
        except HTTPException:
            record = {"record_id": record_id}
    record["requirement_input"] = requirement_input or record.get("requirement_input", []) or []
    record["requirement_spec"] = requirement_spec or record.get("requirement_spec", {}) or {}
    record["persona_reviews"] = persona_reviews or record.get("persona_reviews", {}) or {}
    record["requirement_review_history"] = requirement_review_history or record.get("requirement_review_history", {}) or {}
    record["test_design_spec"] = test_design_spec or record.get("test_design_spec", {}) or {}
    record["product_profile"] = product_profile or record.get("product_profile", {}) or {}
    record["test_environments"] = test_environments or record.get("test_environments", {}) or {}
    if test_case_spec is not None:
        record["test_case_spec"] = test_case_spec
    else:
        record.setdefault("test_case_spec", {})
    if warnings:
        merged = record.get("warnings", {}) if isinstance(record.get("warnings"), dict) else {}
        merged.update(warnings)
        record["warnings"] = merged
    else:
        record.setdefault("warnings", {})
    if trace:
        merged = record.get("trace", {}) if isinstance(record.get("trace"), dict) else {}
        merged.update(trace)
        record["trace"] = merged
    else:
        record.setdefault("trace", {})
    if status:
        record["status"] = status
    return _write_design_history_record(record)


def _list_design_history_records() -> List[Dict[str, Any]]:
    items: List[Dict[str, Any]] = []
    for path in sorted(design_history_dir.glob("*/record.json"), reverse=True):
        try:
            record = json.loads(path.read_text(encoding="utf-8"))
            items.append(_history_summary(record))
        except Exception:
            continue
    items.sort(key=lambda x: x.get("updated_at", ""), reverse=True)
    return items


@app.get("/")
def index():
    return FileResponse(os.path.join(static_dir, "index.html"))


@app.get("/api/design-history")
def list_design_history():
    return {"items": _list_design_history_records()}


@app.get("/api/design-history/{record_id}")
def get_design_history(record_id: str):
    record = _load_design_history_record(record_id)
    return {
        **record,
        "summary": _history_summary(record),
        "document_url": _history_document_url(record_id),
    }


@app.get("/api/design-history/{record_id}/document")
def get_design_history_document(record_id: str):
    record = _load_design_history_record(record_id)
    html_path = _history_document_path(record_id)
    if not html_path.exists():
        html_path.write_text(_render_design_document_html(record), encoding="utf-8")
    return FileResponse(html_path, media_type="text/html", filename=f"{record.get('title', record_id)}-测试设计文档.html")


def _build_effective_rag_context(requirements: List[str], manual_rag_context: str) -> tuple[str, List[Dict[str, Any]]]:
    query = "\n".join(requirements).strip()
    hits = rag_store.retrieve(query, top_k=6)
    auto_context = ""
    if hits:
        lines = []
        for i, h in enumerate(hits, start=1):
            text = (h.get("text", "") or "").strip()
            excerpt = text[:700]
            lines.append(
                f"[RAG-{i}] doc={h.get('filename')} chunk={h.get('chunk_id')} score={h.get('score')}\n{excerpt}"
            )
        auto_context = "\n\n".join(lines)

    parts = [manual_rag_context.strip(), auto_context.strip()]
    merged = "\n\n".join([p for p in parts if p])
    return merged, hits


@app.get("/api/bugzilla/settings")
def bugzilla_get_settings():
    return bugzilla_settings.public()


@app.post("/api/bugzilla/settings")
def bugzilla_save_settings(payload: BugzillaSettingsRequest):
    try:
        return bugzilla_settings.save(payload.url, payload.api_key, payload.clear_key)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@app.post("/api/bugzilla/fetch")
async def bugzilla_fetch(payload: BugzillaFetchRequest):
    cfg = bugzilla_settings.read()
    try:
        bugs = await asyncio.to_thread(fetch_bugs, payload.ids, cfg, payload.excluded_authors)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return {
        "bugs": bugs,
        "bug_ids": [b["id"] for b in bugs if b.get("status") == "ok"],
        "requirement_text": bugs_to_requirement_text(bugs),
    }


@app.get("/api/settings/llm")
def llm_get_settings():
    return LLMSettings().public()


@app.post("/api/settings/llm")
def llm_save_settings(payload: LLMSettingsRequest):
    store = LLMSettings()
    try:
        cfg = store.merged(payload.api_base, payload.model, payload.api_key, payload.clear_key, payload.timeout)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return store.save(cfg)


@app.post("/api/settings/llm/test")
def llm_test_settings(payload: LLMSettingsRequest):
    """用表单里的值（未保存也可以）发一次最小请求，验证地址、key 和模型是否可用。"""
    try:
        cfg = LLMSettings().merged(payload.api_base, payload.model, payload.api_key, payload.clear_key, payload.timeout)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    client = LLMClient(cfg)
    if not client.enabled:
        raise HTTPException(status_code=400, detail="请填写接口地址、API Key 和模型名称。")
    started = time.time()
    try:
        reply = client.generate("You are a connectivity check.", "Reply with the single word: OK", temperature=0, max_tokens=8)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=str(exc))
    return {"ok": True, "reply": reply[:80], "latency_ms": int((time.time() - started) * 1000)}


@app.get("/api/rag/docs")
def rag_docs():
    return {"docs": rag_store.list_docs()}


@app.post("/api/rag/upload")
async def rag_upload(payload: RAGUploadBase64Request):
    filename = (payload.filename or "").strip()
    if not filename:
        raise HTTPException(status_code=400, detail="filename不能为空")
    try:
        content = base64.b64decode(payload.content_base64)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"base64解析失败: {exc}")
    if not content:
        raise HTTPException(status_code=400, detail="文件内容为空")
    try:
        meta = rag_store.ingest_file(filename, content)
        return {"uploaded": meta}
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"文档入库失败: {exc}")


@app.post("/api/chat")
async def chat(payload: ChatRequest):
    message = (payload.message or "").strip()
    if not message:
        raise HTTPException(status_code=400, detail="message不能为空")

    client = LLMClient()
    if not client.enabled:
        return {"reply": "大模型未配置。请先在“系统设置 → 大模型”中填写接口地址、API Key 和模型名称。"}

    history_lines: List[str] = []
    for turn in (payload.history or [])[-8:]:
        role = "用户" if turn.role == "user" else "助手"
        txt = (turn.content or "").strip()
        if txt:
            history_lines.append(f"{role}: {txt}")
    history_text = "\n".join(history_lines)
    user_prompt = (
        (f"对话历史:\n{history_text}\n\n" if history_text else "")
        + f"用户最新问题:\n{message}\n\n请直接回答，简洁、准确、可执行。"
    )
    try:
        reply = client.generate(
            system_prompt="You are a helpful website chatbot. Be concise, accurate, and actionable.",
            user_prompt=user_prompt,
            temperature=0.3,
        )
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"chat失败: {exc}")
    return {"reply": reply}


def _collect_review_issues(persona_reviews: Dict[str, Any]) -> List[str]:
    issues: List[str] = []
    for persona in ["spec", "carrier", "ux"]:
        for item in (persona_reviews.get(persona, {}) or {}).get("reviews", []) or []:
            req_id = item.get("req_id", "")
            for issue in item.get("issues", []) or []:
                issues.append(f"[{persona}|{req_id}] {issue}")
    return issues


MAX_USER_QUESTIONS = 5
_PRIORITY_RANK = {"P0": 0, "P1": 1, "P2": 2}
_PERSONA_LABELS = {"spec": "规范", "carrier": "运营商", "ux": "体验"}


def _question_key(text: str) -> str:
    return re.sub(r"[\s，。、,.；;：:？?！!“”\"'（）()]+", "", str(text or "")).lower()


def _triage_open_questions(persona_reviews: Dict[str, Any], answered: Any = None) -> tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """把三位评审提出的问题分成两组：
    - 需要用户确认：needs_user 且优先级最高的少数几条（最多 MAX_USER_QUESTIONS 条）
    - AI 自行确认：其余问题，直接采用评审给出的默认答案
    已经回答过的问题、三位评审重复提出的问题会被去掉。"""
    done = {_question_key(q.get("question")) for q in (answered or []) if isinstance(q, dict)}
    merged: Dict[str, Dict[str, Any]] = {}
    for persona in ["spec", "carrier", "ux"]:
        for item in (persona_reviews.get(persona, {}) or {}).get("reviews", []) or []:
            req_id = str(item.get("req_id", ""))
            details = item.get("question_details")
            if not isinstance(details, list):
                details = [{"question": q, "priority": "P1", "needs_user": True, "ai_answer": "", "basis": ""}
                           for q in item.get("open_questions", []) or [] if isinstance(q, str)]
            for d in details:
                if not isinstance(d, dict) or not str(d.get("question") or "").strip():
                    continue
                key = _question_key(d["question"])
                if not key or key in done:
                    continue
                cur = merged.get(key)
                cand = {
                    "req_id": req_id,
                    "question": str(d["question"]).strip(),
                    "priority": d.get("priority") if d.get("priority") in _PRIORITY_RANK else "P1",
                    "needs_user": bool(d.get("needs_user", True)),
                    "ai_answer": str(d.get("ai_answer") or "").strip(),
                    "basis": str(d.get("basis") or "").strip(),
                    "personas": [_PERSONA_LABELS[persona]],
                }
                if cur is None:
                    merged[key] = cand
                    continue
                # 多位评审都提到的问题：合并来源，取更高优先级，任一评审认为需要用户确认就保留
                cur["personas"] = list(dict.fromkeys(cur["personas"] + cand["personas"]))
                if _PRIORITY_RANK[cand["priority"]] < _PRIORITY_RANK[cur["priority"]]:
                    cur["priority"] = cand["priority"]
                cur["needs_user"] = cur["needs_user"] or cand["needs_user"]
                if not cur["ai_answer"] and cand["ai_answer"]:
                    cur["ai_answer"], cur["basis"] = cand["ai_answer"], cand["basis"]

    ordered = sorted(merged.values(), key=lambda q: (_PRIORITY_RANK[q["priority"]], -len(q["personas"])))
    user_qs: List[Dict[str, Any]] = []
    auto_qs: List[Dict[str, Any]] = []
    for q in ordered:
        must_ask = q["needs_user"] and (q["priority"] in ("P0", "P1") or not q["ai_answer"])
        if must_ask and len(user_qs) < MAX_USER_QUESTIONS:
            user_qs.append(q)
        else:
            if not q["ai_answer"]:
                q["ai_answer"] = "暂按需求原文理解，测试设计中标注为待确认风险"
            auto_qs.append(q)
    return user_qs, auto_qs


def _collect_open_questions(persona_reviews: Dict[str, Any]) -> List[Dict[str, Any]]:
    return _triage_open_questions(persona_reviews)[0]


async def _emit_hitl_event(job: Dict[str, Any], ev: Dict[str, Any]):
    """Broadcast HITL events to internal control queue and all SSE subscribers."""
    control_q: asyncio.Queue = job["queue"]
    await control_q.put(ev)
    job.setdefault("recent_events", []).append(ev)
    if len(job["recent_events"]) > 80:
        job["recent_events"] = job["recent_events"][-80:]

    stale: List[asyncio.Queue] = []
    for sub_q in job.get("subscribers", []) or []:
        try:
            sub_q.put_nowait(ev)
        except Exception:
            stale.append(sub_q)
    if stale:
        job["subscribers"] = [q for q in (job.get("subscribers", []) or []) if q not in stale]


async def _emit_stage_event(job: Dict[str, Any], ev: Dict[str, Any]):
    """Broadcast stage events to all subscribers and keep a short replay window."""
    control_q: asyncio.Queue = job["queue"]
    await control_q.put(ev)
    job.setdefault("recent_events", []).append(ev)
    if len(job["recent_events"]) > 120:
        job["recent_events"] = job["recent_events"][-120:]

    stale: List[asyncio.Queue] = []
    for sub_q in job.get("subscribers", []) or []:
        try:
            sub_q.put_nowait(ev)
        except Exception:
            stale.append(sub_q)
    if stale:
        job["subscribers"] = [q for q in (job.get("subscribers", []) or []) if q not in stale]


async def _run_requirement_hitl_job(job_id: str):
    job = requirement_hitl_jobs[job_id]
    queue: asyncio.Queue = job["queue"]
    reqs: List[str] = job["requirements"]
    rag_context: str = job["rag_context"]
    framework_capability_catalog: str = job["framework_capability_catalog"]
    product_profile: Dict[str, Any] = job.get("product_profile") or {}
    test_environments: Dict[str, Any] = job.get("test_environments") or {}

    try:
        for round_no in range(1, 4):
            effective_rag_context, rag_hits = _build_effective_rag_context(reqs, rag_context)
            shared = {
                "input_requirements": reqs,
                "rag_context": effective_rag_context,
                "framework_capability_catalog": framework_capability_catalog,
                "product_profile": product_profile,
                "test_environments": test_environments,
                "capabilities": DEFAULT_CAPABILITIES,
                "action_vocabulary": DEFAULT_ACTION_VOCABULARY,
                "assertion_vocabulary": DEFAULT_ASSERTION_VOCABULARY,
                "llm_client": LLMClient(),
                "warnings": [],
                "trace": [],
                "review_feedback": job.get("review_feedback", ""),
                "event_queue": queue,
            }
            flow = create_requirement_analysis_flow()
            await flow.run_async(shared)

            requirement_spec = shared.get("requirement_spec", {})
            persona_reviews = shared.get("persona_reviews", {})
            user_questions, auto_confirmed = _triage_open_questions(persona_reviews, job.get("answered_questions", []))
            # AI 自行确认的问题直接记入已确认列表，用户跳过后续轮次直接生成设计时也能带上这些假设
            job["answered_questions"] = _merge_answered_questions(
                job.get("answered_questions", []),
                [{"source_round": round_no, "req_id": q["req_id"], "question": q["question"],
                  "answer": q["ai_answer"], "source": "AI 自行确认"} for q in auto_confirmed],
            )
            job["pending_questions"] = user_questions
            round_summary = {
                "round": round_no,
                "is_final_round": round_no >= 3,
                "issues": _collect_review_issues(persona_reviews),
                "open_questions": user_questions,
                "auto_confirmed": auto_confirmed,
            }
            round_summaries = [item for item in (job.get("round_summaries", []) or []) if item.get("round") != round_no]
            round_summaries.append(round_summary)
            round_summaries.sort(key=lambda item: item.get("round", 0))
            job["round_summaries"] = round_summaries
            review_history = _build_requirement_review_history(job.get("round_summaries", []), job.get("answered_questions", []))
            round_payload = {
                "job_id": job_id,
                "round": round_no,
                "is_final_round": round_no >= 3,
                "requirement_spec": requirement_spec,
                "persona_reviews": persona_reviews,
                "issues": round_summary["issues"],
                "open_questions": round_summary["open_questions"],
                "auto_confirmed": auto_confirmed,
                "requirement_review_history": review_history,
                "warnings": shared.get("warnings", []),
                "trace": shared.get("trace", []),
                "rag": {"hits": rag_hits, "context_used": bool(effective_rag_context)},
            }
            job["last_payload"] = round_payload
            job["round"] = round_no
            job["status"] = "round_ready"
            await _emit_hitl_event(job, {"type": "round_ready", "payload": round_payload})

            if round_no >= 3:
                job["status"] = "completed"
                await _emit_hitl_event(job, {"type": "completed", "payload": round_payload})
                return

            job["status"] = "waiting_answers"
            await _emit_hitl_event(
                job,
                {
                    "type": "waiting_answers",
                    "payload": {"job_id": job_id, "round": round_no, "open_questions": round_payload["open_questions"]},
                }
            )
            await job["answer_event"].wait()
            job["answer_event"].clear()
            answers = {k: v for k, v in (job.get("pending_answers", {}) or {}).items() if str(v).strip()}
            # 用户没填的问题采用 AI 建议答案
            accepted = []
            for q in job.get("pending_questions", []) or []:
                key = f"{q['req_id']}::{q['question']}"
                if key not in answers and q.get("ai_answer"):
                    accepted.append({"source_round": round_no, "req_id": q["req_id"], "question": q["question"],
                                     "answer": q["ai_answer"], "source": "采纳 AI 建议"})
            job["answered_questions"] = _merge_answered_questions(
                job.get("answered_questions", []),
                _parse_answer_map(answers, round_no) + accepted,
            )

            fb_parts: List[str] = []
            issues = _collect_review_issues(persona_reviews)
            if issues:
                fb_parts.append("【上轮评审issues回灌】" + " | ".join(issues))
            answer_lines = [f"{k}: {v}" for k, v in answers.items() if str(v).strip()]
            if answer_lines:
                fb_parts.append("【用户补充回答】" + " | ".join(answer_lines))
            ai_lines = [f"{q['req_id']}::{q['question']}: {q['answer']}" for q in job.get("answered_questions", [])
                        if q.get("source") != "用户"]
            if ai_lines:
                fb_parts.append("【AI 自行确认 / 采纳的默认假设，可直接写入需求】" + " | ".join(ai_lines))
            job["review_feedback"] = "\n".join(fb_parts).strip()
            job["pending_answers"] = {}

    except Exception as exc:
        job["status"] = "failed"
        job["error"] = str(exc)
        await _emit_hitl_event(job, {"type": "error", "payload": {"job_id": job_id, "error": str(exc)}})


async def _wait_hitl_round(job_id: str, timeout_s: int = 180):
    job = requirement_hitl_jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="job不存在")
    queue: asyncio.Queue = job["queue"]
    try:
        while True:
            ev = await asyncio.wait_for(queue.get(), timeout=timeout_s)
            et = ev.get("type")
            if et in {"round_ready", "completed"}:
                return ev.get("payload", {})
            if et == "error":
                msg = ev.get("payload", {}).get("error", "unknown error")
                raise HTTPException(status_code=500, detail=f"HITL执行失败: {msg}")
    except asyncio.TimeoutError:
        raise HTTPException(status_code=504, detail="等待HITL结果超时")


@app.post("/api/hitl/requirements/start")
async def hitl_requirements_start(payload: RequirementHITLStartRequest):
    reqs = [r.strip() for r in payload.requirements if r.strip()]
    if not reqs:
        raise HTTPException(status_code=400, detail="requirements不能为空")

    job_id = uuid.uuid4().hex[:12]
    queue: asyncio.Queue = asyncio.Queue()
    answer_event = asyncio.Event()
    framework_capability_catalog = payload.framework_capability_catalog.strip()

    requirement_hitl_jobs[job_id] = {
        "job_id": job_id,
        "status": "starting",
        "round": 0,
        "requirements": reqs,
        "rag_context": payload.rag_context.strip(),
        "framework_capability_catalog": framework_capability_catalog,
        "product_profile": payload.product_profile or {},
        "test_environments": payload.test_environments or {},
        "review_feedback": "",
        "round_summaries": [],
        "answered_questions": [],
        "pending_answers": {},
        "answer_event": answer_event,
        "queue": queue,
        "subscribers": [],
        "recent_events": [],
        "last_payload": None,
        "error": "",
    }
    requirement_hitl_jobs[job_id]["task"] = asyncio.create_task(_run_requirement_hitl_job(job_id))
    first = await _wait_hitl_round(job_id)
    return first


@app.post("/api/hitl/requirements/{job_id}/next")
async def hitl_requirements_next(job_id: str, payload: RequirementHITLNextRequest):
    job = requirement_hitl_jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="job不存在")
    if job.get("status") not in {"waiting_answers", "round_ready"}:
        if job.get("status") == "completed":
            return job.get("last_payload") or {"job_id": job_id, "is_final_round": True}
        raise HTTPException(status_code=409, detail=f"当前状态不允许next: {job.get('status')}")
    job["pending_answers"] = payload.open_question_answers or {}
    job["answer_event"].set()
    nxt = await _wait_hitl_round(job_id)
    return nxt


@app.get("/api/hitl/requirements/{job_id}/state")
def hitl_requirements_state(job_id: str):
    job = requirement_hitl_jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="job不存在")
    return {
        "job_id": job_id,
        "status": job.get("status"),
        "round": job.get("round", 0),
        "last_payload": job.get("last_payload"),
        "error": job.get("error", ""),
    }


@app.get("/api/hitl/requirements/{job_id}/stream")
async def hitl_requirements_stream(job_id: str):
    job = requirement_hitl_jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="job不存在")

    sub_q: asyncio.Queue = asyncio.Queue()
    subscribers = job.setdefault("subscribers", [])
    subscribers.append(sub_q)

    async def event_generator():
        try:
            init_payload = {
                "job_id": job_id,
                "status": job.get("status"),
                "round": job.get("round", 0),
                "error": job.get("error", ""),
            }
            yield f"data: {json.dumps({'type': 'state', 'payload': init_payload}, ensure_ascii=False)}\n\n"

            for ev in job.get("recent_events", [])[-20:]:
                yield f"data: {json.dumps(ev, ensure_ascii=False)}\n\n"

            while True:
                ev = await sub_q.get()
                yield f"data: {json.dumps(ev, ensure_ascii=False)}\n\n"
                if ev.get("type") in {"completed", "error"}:
                    break
        finally:
            if sub_q in (job.get("subscribers", []) or []):
                job["subscribers"] = [q for q in job.get("subscribers", []) if q is not sub_q]

    return StreamingResponse(event_generator(), media_type="text/event-stream")


async def _run_stage_job(job_id: str):
    job = stage_jobs[job_id]
    stage = job["stage"]
    payload: StageGenerateRequest = job["payload"]
    queue: asyncio.Queue = job["queue"]

    await _emit_stage_event(job, {"type": "job_started", "payload": {"job_id": job_id, "stage": stage}})

    try:
        if stage == "design":
            shared = {
                "input_requirements": payload.requirement_input or [],
                "requirement_spec": payload.requirement_spec,
                "persona_reviews": payload.persona_reviews or {},
                "requirement_review_history": payload.requirement_review_history or {},
                "product_profile": payload.product_profile or {},
                "test_environments": payload.test_environments or {},
                "llm_client": LLMClient(),
                "warnings": [],
                "trace": [],
                "event_queue": queue,
            }
            flow = create_test_design_flow()
            await flow.run_async(shared)
            history_record = _upsert_design_history_record(
                record_id=payload.history_record_id or "",
                requirement_input=payload.requirement_input or [],
                requirement_spec=payload.requirement_spec or {},
                persona_reviews=payload.persona_reviews or {},
                requirement_review_history=payload.requirement_review_history or {},
                test_design_spec=shared.get("test_design_spec", {}),
                product_profile=payload.product_profile or {},
                test_environments=payload.test_environments or {},
                warnings={"design": shared.get("warnings", [])},
                trace={"design": shared.get("trace", [])},
                status="design_completed",
            )
            result = {
                "test_design_spec": shared.get("test_design_spec", {}),
                "warnings": shared.get("warnings", []),
                "trace": shared.get("trace", []),
                "design_history": _history_summary(history_record),
            }
        elif stage == "testcases":
            if not payload.test_design_spec:
                raise HTTPException(status_code=400, detail="缺少test_design_spec")
            shared = {
                "input_requirements": payload.requirement_input or [],
                "requirement_spec": payload.requirement_spec,
                "test_design_spec": payload.test_design_spec,
                "requirement_review_history": payload.requirement_review_history or {},
                "product_profile": payload.product_profile or {},
                "test_environments": payload.test_environments or {},
                "action_vocabulary": payload.action_vocabulary or DEFAULT_ACTION_VOCABULARY,
                "llm_client": LLMClient(),
                "warnings": [],
                "trace": [],
                "event_queue": queue,
            }
            flow = create_test_case_flow()
            await flow.run_async(shared)
            history_record = _upsert_design_history_record(
                record_id=payload.history_record_id or "",
                requirement_input=payload.requirement_input or [],
                requirement_spec=payload.requirement_spec or {},
                persona_reviews=payload.persona_reviews or {},
                requirement_review_history=payload.requirement_review_history or {},
                test_design_spec=payload.test_design_spec or {},
                product_profile=payload.product_profile or {},
                test_environments=payload.test_environments or {},
                test_case_spec=shared.get("test_case_spec", {}),
                warnings={"testcases": shared.get("warnings", [])},
                trace={"testcases": shared.get("trace", [])},
                status="testcases_completed",
            )
            result = {
                "test_case_spec": shared.get("test_case_spec", {}),
                "warnings": shared.get("warnings", []),
                "trace": shared.get("trace", []),
                "design_history": _history_summary(history_record),
            }
        elif stage == "exec_cases":
            if not payload.test_case_spec:
                raise HTTPException(status_code=400, detail="缺少test_case_spec")
            shared = {
                "requirement_spec": payload.requirement_spec or {},
                "test_case_spec": payload.test_case_spec,
                "product_profile": payload.product_profile or {},
                "test_environments": payload.test_environments or {},
                "bug_ids": payload.bug_ids or [],
                "route_overrides": payload.route_overrides or {},
                "llm_client": LLMClient(),
                "warnings": [],
                "trace": [],
                "event_queue": queue,
            }
            flow = create_exec_case_flow()
            await flow.run_async(shared)
            executable_cases = shared.get("executable_cases", {})
            history = None
            if payload.history_record_id and _valid_history_record_id(payload.history_record_id):
                try:
                    record = _load_design_history_record(payload.history_record_id)
                    record["executable_cases"] = executable_cases
                    record["status"] = "exec_cases_completed"
                    history = _history_summary(_write_design_history_record(record))
                except HTTPException:
                    history = None
            result = {
                "executable_cases": executable_cases,
                "warnings": shared.get("warnings", []),
                "trace": shared.get("trace", []),
                "design_history": history,
            }
        else:
            raise RuntimeError(f"不支持的stage: {stage}")

        job["status"] = "completed"
        job["result"] = result
        await _emit_stage_event(
            job,
            {"type": "completed", "payload": {"job_id": job_id, "stage": stage, "result": result}},
        )
    except Exception as exc:
        job["status"] = "failed"
        detail = {
            "error": str(exc),
            "failed_node": (locals().get("shared") or {}).get("current_node"),
            "trace": (locals().get("shared") or {}).get("trace", []),
            "warnings": (locals().get("shared") or {}).get("warnings", []),
        }
        job["error"] = detail
        await _emit_stage_event(
            job,
            {"type": "error", "payload": {"job_id": job_id, "stage": stage, **detail}},
        )


@app.post("/api/stage/{stage}/start")
async def stage_start(stage: str, payload: StageGenerateRequest):
    stage_name = (stage or "").strip().lower()
    if stage_name not in {"design", "testcases", "exec_cases"}:
        raise HTTPException(status_code=400, detail="stage必须是design/testcases/exec_cases")

    if stage_name == "design" and not payload.requirement_spec:
        raise HTTPException(status_code=400, detail="缺少requirement_spec")
    if stage_name == "testcases" and not payload.test_design_spec:
        raise HTTPException(status_code=400, detail="缺少test_design_spec")
    if stage_name == "exec_cases" and not payload.test_case_spec:
        raise HTTPException(status_code=400, detail="缺少test_case_spec")

    job_id = uuid.uuid4().hex[:12]
    q: asyncio.Queue = asyncio.Queue()
    stage_jobs[job_id] = {
        "job_id": job_id,
        "stage": stage_name,
        "status": "running",
        "payload": payload,
        "queue": q,
        "subscribers": [],
        "recent_events": [],
        "result": None,
        "error": None,
    }
    stage_jobs[job_id]["task"] = asyncio.create_task(_run_stage_job(job_id))
    return {"job_id": job_id, "stage": stage_name, "status": "running"}


@app.get("/api/stage/{job_id}/state")
def stage_state(job_id: str):
    job = stage_jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="job不存在")
    return {
        "job_id": job_id,
        "stage": job.get("stage"),
        "status": job.get("status"),
        "result": job.get("result"),
        "error": job.get("error"),
    }


@app.get("/api/stage/{job_id}/stream")
async def stage_stream(job_id: str):
    job = stage_jobs.get(job_id)
    if not job:
        raise HTTPException(status_code=404, detail="job不存在")
    queue: asyncio.Queue = job["queue"]

    async def event_generator():
        init_payload = {
            "job_id": job_id,
            "stage": job.get("stage"),
            "status": job.get("status"),
            "error": job.get("error"),
        }
        yield f"data: {json.dumps({'type': 'state', 'payload': init_payload}, ensure_ascii=False)}\n\n"
        for ev in job.get("recent_events", [])[-40:]:
            yield f"data: {json.dumps(ev, ensure_ascii=False)}\n\n"

        while True:
            ev = await queue.get()
            yield f"data: {json.dumps(ev, ensure_ascii=False)}\n\n"
            if ev.get("type") in {"completed", "error"}:
                break

    return StreamingResponse(event_generator(), media_type="text/event-stream")


@app.get("/api/settings/prompts")
def get_prompts():
    data: Dict[str, str] = {}
    for name in sorted(ALLOWED_PROMPT_FILES):
        p = prompt_path(name)
        if p.exists():
            data[name] = p.read_text(encoding="utf-8")
    return {"prompts": data}


def _flow_to_graph(flow_obj: Any) -> Dict[str, Any]:
    start = getattr(flow_obj, "start_node", None)
    if not start:
        return {"nodes": [], "links": [], "mermaid": "flowchart LR\n", "start_node_id": ""}

    node_id_map: Dict[int, str] = {}
    nodes: List[Dict[str, Any]] = []
    links: List[Dict[str, Any]] = []
    queue: List[Any] = [start]
    visited: set[int] = set()
    edge_seen: set[tuple[str, str, str]] = set()
    counter = 1

    while queue:
        node = queue.pop(0)
        obj_id = id(node)
        if obj_id in visited:
            continue
        visited.add(obj_id)
        if obj_id not in node_id_map:
            node_id_map[obj_id] = f"N{counter}"
            counter += 1
        nid = node_id_map[obj_id]
        nodes.append({"id": nid, "label": type(node).__name__})
        for action, nxt in (getattr(node, "successors", {}) or {}).items():
            nxt_id = id(nxt)
            if nxt_id not in node_id_map:
                node_id_map[nxt_id] = f"N{counter}"
                counter += 1
            tid = node_id_map[nxt_id]
            action_str = str(action or "default")
            ekey = (nid, tid, action_str)
            if ekey not in edge_seen:
                edge_seen.add(ekey)
                links.append({"source": nid, "target": tid, "action": action_str})
            queue.append(nxt)

    mermaid_lines = ["flowchart LR"]
    for n in nodes:
        mermaid_lines.append(f'    {n["id"]}["{n["label"]}"]')
    for e in links:
        mermaid_lines.append(f'    {e["source"]} -- "{e["action"]}" --> {e["target"]}')
    start_node_id = node_id_map.get(id(start), "")
    return {"nodes": nodes, "links": links, "mermaid": "\n".join(mermaid_lines), "start_node_id": start_node_id}


def _prompt_role_mapping() -> List[Dict[str, Any]]:
    role_map = [
        {"role": "Requirement Parser", "node": "ReqParseNode", "prompt": "req_parse.txt"},
        {"role": "Spec Lawyer", "node": "TriPersonaReviewNode", "prompt": "review_spec_lawyer.txt"},
        {"role": "Carrier Reviewer", "node": "TriPersonaReviewNode", "prompt": "review_carrier.txt"},
        {"role": "UX Advocate", "node": "TriPersonaReviewNode", "prompt": "review_ux_advocate.txt"},
        {"role": "Synthesis Arbiter", "node": "SynthesisNode", "prompt": "synthesis_arbiter.txt"},
        {"role": "Test Designer", "node": "TestDesignNode", "prompt": "test_designer.txt"},
        {"role": "Integrated Matrix Designer", "node": "TestDesignNode", "prompt": "integrated_matrix_designer.txt"},
        {"role": "Test Case Generator", "node": "TestCaseBatchGenNode", "prompt": "test_case_generator.txt"},
        {"role": "Artemis Case Writer", "node": "ExecutableCaseNode", "prompt": "exec_case_artemis.txt"},
        {"role": "AT Agent Case Writer", "node": "ExecutableCaseNode", "prompt": "exec_case_at_agent.txt"},
        {"role": "Amarisoft Case Writer", "node": "ExecutableCaseNode", "prompt": "exec_case_amarisoft.txt"},
        {"role": "Protocol Flow Requirement Writer", "node": "ExecutableCaseNode", "prompt": "exec_case_protocol.txt"},
    ]
    out: List[Dict[str, Any]] = []
    for item in role_map:
        p = prompt_path(item["prompt"])
        first_line = ""
        if p.exists():
            for ln in p.read_text(encoding="utf-8").splitlines():
                raw = ln.strip()
                if raw:
                    first_line = raw[:120]
                    break
        out.append(
            {
                **item,
                "exists": p.exists(),
                "first_line": first_line,
            }
        )
    return out


@app.get("/api/settings/flow-visualization")
def get_flow_visualization():
    flow = create_system_test_flow()
    graph = _flow_to_graph(flow)
    return {
        "flow_name": "SystemTestFlow",
        "mermaid": graph["mermaid"],
        "graph": {"nodes": graph["nodes"], "links": graph["links"], "start_node_id": graph.get("start_node_id", "")},
        "prompt_roles": _prompt_role_mapping(),
    }


@app.post("/api/settings/prompts")
def save_prompts(payload: PromptSettingsRequest):
    for name, content in payload.prompts.items():
        if name not in ALLOWED_PROMPT_FILES:
            raise HTTPException(status_code=400, detail=f"非法模板名: {name}")
        PROMPT_OVERRIDE_DIR.mkdir(parents=True, exist_ok=True)
        (PROMPT_OVERRIDE_DIR / name).write_text(content, encoding="utf-8")
    return {"message": "saved", "count": len(payload.prompts)}


@app.post("/api/export/testcases.xlsx")
def export_testcases_excel(payload: ExportTestCasesRequest):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"缺少openpyxl依赖: {exc}")

    requirement_input = payload.requirement_input or []
    requirement_spec = payload.requirement_spec or {}
    persona_reviews = payload.persona_reviews or {}
    requirement_review_history = payload.requirement_review_history or {}
    test_design_spec = payload.test_design_spec or {}
    test_case_spec = payload.test_case_spec or {}
    testcases = test_case_spec.get("testcases", [])

    def _list_to_text(v: Any) -> str:
        if isinstance(v, list):
            return "\n".join([str(x) for x in v])
        if v is None:
            return ""
        return str(v)

    def _safe_json(v: Any) -> str:
        try:
            return json.dumps(v, ensure_ascii=False)
        except Exception:
            return str(v)

    def _kv_text(v: Any) -> str:
        if isinstance(v, dict):
            return "\n".join(f"{k}: {_list_to_text(x) if isinstance(x, list) else x}" for k, x in v.items())
        return _list_to_text(v)

    def _numbered(v: Any) -> str:
        items = [str(x) for x in v] if isinstance(v, list) else ([str(v)] if v else [])
        return items[0] if len(items) == 1 else "\n".join(f"{i}. {x}" for i, x in enumerate(items, 1))

    def _sheet(title: str, headers: List[tuple], rows: List[List[Any]], first: bool = False):
        """headers: [(列名, 列宽)]。表头加粗、冻结首行、自动换行。"""
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append([h for h, _ in headers])
        for r in rows:
            ws.append(r)
        for idx, (_, width) in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        return ws

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    persona_names = {"spec": "规范评审", "carrier": "运营商评审", "ux": "体验评审"}

    case_rows = []
    for tc in testcases:
        if not isinstance(tc, dict):
            continue
        env_text, risk_text = _format_test_environment_cell(tc.get("test_environment"))
        case_rows.append([
            tc.get("tc_id", ""), tc.get("title", ""), tc.get("objective_id", ""),
            _numbered(tc.get("preconditions", [])), _numbered(tc.get("expected", [])), _numbered(tc.get("pass_fail", [])),
            env_text, risk_text, "、".join(str(x) for x in tc.get("tags", []) or []),
        ])
    _sheet("测试用例", [("用例编号", 14), ("用例标题", 40), ("测试目标", 12), ("前置条件", 40), ("预期结果", 50),
                        ("通过标准", 50), ("测试环境", 18), ("环境风险", 30), ("标签", 18)], case_rows, first=True)

    matrix_rows = []
    for row in test_design_spec.get("integrated_matrix", []) if isinstance(test_design_spec, dict) else []:
        if not isinstance(row, dict):
            continue
        env_text, risk_text = _format_test_environment_cell(row.get("test_environment"))
        matrix_rows.append([
            row.get("row_id", ""), row.get("req_id", ""), row.get("objective_id", ""), row.get("scenario", ""),
            _kv_text(row.get("key_configuration", {})), _numbered(row.get("pass_criteria", [])), env_text, risk_text,
        ])
    _sheet("测试设计矩阵", [("矩阵行", 12), ("需求", 12), ("测试目标", 12), ("测试场景", 45), ("关键配置", 40),
                          ("通过标准", 45), ("测试环境", 18), ("环境风险", 30)], matrix_rows)

    obj_rows = []
    for o in test_design_spec.get("objectives", []) if isinstance(test_design_spec, dict) else []:
        if isinstance(o, dict):
            obj_rows.append([
                o.get("objective_id", ""), _list_to_text(o.get("linked_reqs", [])), o.get("goal", ""),
                _numbered(o.get("success_criteria", [])), _numbered(o.get("evidence", [])), o.get("priority", ""), o.get("risk_notes", ""),
            ])
    _sheet("测试设计目标", [("测试目标", 12), ("关联需求", 14), ("目标", 45), ("成功标准", 45), ("证据", 35),
                          ("优先级", 8), ("风险说明", 35)], obj_rows)

    req_rows = []
    for r in requirement_spec.get("final_requirements", []) if isinstance(requirement_spec, dict) else []:
        if isinstance(r, dict):
            req_rows.append([
                r.get("req_id", ""), r.get("title", ""), r.get("priority", ""), _list_to_text(r.get("rat_scope", [])),
                _numbered((r.get("acceptance") or {}).get("pass_fail", [])), _numbered((r.get("acceptance") or {}).get("kpi", [])),
            ])
    _sheet("测试需求", [("需求编号", 12), ("需求标题", 45), ("优先级", 8), ("制式范围", 14), ("通过/失败判定", 50),
                      ("KPI", 40)], req_rows)

    answer_rows = [
        [item.get("source_round", ""), item.get("req_id", ""), item.get("question", ""), item.get("answer", ""),
         item.get("source", "用户")]
        for item in (requirement_review_history.get("answered_questions", []) if isinstance(requirement_review_history, dict) else [])
        if isinstance(item, dict)
    ]
    if answer_rows:
        _sheet("问题确认记录", [("轮次", 8), ("需求", 12), ("问题", 50), ("确认结果", 50), ("确认方", 14)], answer_rows)

    review_rows = []
    for persona_key in ["spec", "carrier", "ux"]:
        revs = ((persona_reviews.get(persona_key) or {}).get("reviews", []) if isinstance(persona_reviews, dict) else [])
        for item in revs:
            if isinstance(item, dict):
                review_rows.append([persona_names[persona_key], item.get("req_id", ""), _numbered(item.get("issues", [])),
                                    item.get("rewrite_suggestion", "")])
    if review_rows:
        _sheet("三人评审", [("评审角色", 12), ("需求", 12), ("发现的问题", 60), ("改写建议", 60)], review_rows)

    executable_cases = payload.executable_cases or {}
    trace_rows = executable_cases.get("traceability") or []
    if trace_rows:
        _sheet("追溯矩阵", [("Bug ID", 12), ("需求 ID", 12), ("需求标题", 35), ("测试目标", 12), ("矩阵行", 12),
                          ("用例 ID", 14), ("用例标题", 35), ("测试环境", 16), ("验收方案", 22), ("可执行 case", 16)], [
            [_list_to_text(r.get("bug_ids")), _list_to_text(r.get("req_ids")), _list_to_text(r.get("req_titles")),
             r.get("objective_id", ""), _list_to_text(r.get("row_ids")), r.get("tc_id", ""), r.get("tc_title", ""),
             r.get("test_environment", ""), r.get("route_label", ""), _list_to_text(r.get("exec_case_ids"))]
            for r in trace_rows
        ])

    if requirement_input:
        _sheet("需求原文", [("序号", 8), ("需求原文", 120)], [[i, txt] for i, txt in enumerate(requirement_input, start=1)])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    headers = {"Content-Disposition": "attachment; filename=testcases.xlsx"}
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8010)
